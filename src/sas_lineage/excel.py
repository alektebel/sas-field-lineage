"""
Excel export of a parsed program: one sheet, one row per field.

The sheet is deliberately flat. A workbook with a sheet per table looks tidy
and is useless for the question people actually ask -- "where does this field
come from, and at what point in the pipeline" -- because answering it means
opening every sheet. One sheet with the table, its stage and the upstream flow
on each row answers it with a filter.

Columns
-------
``stage``        pipeline stage of the table (0 = golden source, nothing
                 writes it; each step away from a source adds one)
``libref``       the library, or blank when the name is unresolved
``table``        full ``libref.member`` name as SAS resolves it
``persistent``   yes / no / unknown -- yes when the libref is not WORK, so the
                 table outlives the session
``field``        the field, blank for a table the parser found no fields in
``hop``          the field's position within ``flow`` (1 = a golden input)
``expression``   the SAS expression that builds it
``direct_inputs``the fields it reads directly
``flow``         the whole upstream chain, golden source first
``golden_inputs``the fields at the head of that chain
``source_line``  line in the SAS source
``written_by``   the steps that write the table
``read_by``      the steps that read it
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

from .tables import build_inventory

HEADERS = [
    "stage", "libref", "table", "persistent", "field", "hop", "expression",
    "direct_inputs", "flow", "golden_inputs", "source_line", "written_by", "read_by",
]

_ARROW = " -> "


def _persistent_label(value: Optional[bool]) -> str:
    if value is None:
        return "unknown"
    return "yes" if value else "no"


def build_rows(program) -> List[Dict[str, Any]]:
    """One row per field, plus one row for every table that has no fields.

    A table with no parsed fields (a PROC output, a golden source, a step whose
    body the parser could not read) still belongs in the inventory: leaving it
    out would make the sheet look complete when it is not.
    """
    from .ui.facts import FieldFacts        # local: keeps import cost off the parser

    inventory = build_inventory(program)
    facts = FieldFacts(program)
    rows: List[Dict[str, Any]] = []
    tables_with_fields = set()

    for entry in facts.index.field_entries:
        table = entry["table"]
        info = inventory.get(table)
        tables_with_fields.add(table)
        hops, golden = facts.hops_for(entry["id"])
        # Golden inputs head the flow: they are where the values come from,
        # even though they are not derivations. ``hop`` indexes into this same
        # list so the two columns agree.
        flow = golden + [h["field"] for h in hops if h["field"] not in golden]
        try:
            hop_no = flow.index(entry["id"]) + 1
        except ValueError:
            flow.append(entry["id"])
            hop_no = len(flow)
        node = entry["node"]
        rows.append({
            "stage": info.stage if info else 0,
            "libref": (info.libref if info else None) or "",
            "table": table,
            "persistent": _persistent_label(info.persistent if info else None),
            "field": node.name,
            "hop": hop_no,
            "expression": node.expression or "",
            "direct_inputs": ", ".join(sorted({
                f"{d.table}.{d.name}" if d.table else d.name for d in node.dependencies
            })),
            "flow": _ARROW.join(flow),
            "golden_inputs": ", ".join(golden),
            "source_line": node.source_line or "",
            "written_by": ", ".join(info.written_by) if info else "",
            "read_by": ", ".join(info.read_by) if info else "",
        })

    for name, info in inventory.items():
        if name in tables_with_fields:
            continue
        rows.append({
            "stage": info.stage,
            "libref": info.libref or "",
            "table": name,
            "persistent": _persistent_label(info.persistent),
            "field": "",
            "hop": "",
            "expression": "",
            "direct_inputs": "",
            "flow": "",
            "golden_inputs": "",
            "source_line": "",
            "written_by": ", ".join(info.written_by),
            "read_by": ", ".join(info.read_by),
        })

    rows.sort(key=lambda r: (r["stage"], r["table"], str(r["field"])))
    return rows


def write_workbook(program, path: str, sheet_name: str = "lineage") -> Path:
    """Write the single-sheet lineage export and return the path written."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise RuntimeError(
            "Excel export needs openpyxl (pip install openpyxl)."
        ) from exc

    rows = build_rows(program)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="D04A02")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append([row[h] for h in HEADERS])

    # A long flow string must not set the column width for everything else.
    widths = {"expression": 44, "direct_inputs": 30, "flow": 60,
              "golden_inputs": 30, "written_by": 26, "read_by": 26}
    for idx, header in enumerate(HEADERS, start=1):
        letter = get_column_letter(idx)
        if header in widths:
            ws.column_dimensions[letter].width = widths[header]
        else:
            longest = max([len(header)] + [len(str(r[header])) for r in rows] or [0])
            ws.column_dimensions[letter].width = min(max(longest + 2, 9), 30)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    out = Path(path)
    wb.save(out)
    return out
